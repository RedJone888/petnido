from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

fonts = [
    "Arial Unicode MS",
    "STSong",
    "Songti SC",
    "Hiragino Sans GB",
    "Heiti SC",
    "PingFang SC",
    "GB18030 Bitmap",
    "Noto Sans CJK SC",
    "Microsoft YaHei",
    "SimSun",
]

doc = Document()
for font in fonts:
    p = doc.add_paragraph()
    r = p.add_run(f"{font}: 宠物照护需求与服务撮合平台 项目计划书 中文测试")
    r.font.name = font
    r.font.size = Pt(14)
    rfonts = r._element.get_or_add_rPr().rFonts
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), font)

doc.save(".codex-qa/cjk-font-test.docx")
